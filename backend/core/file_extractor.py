import io
from pathlib import Path


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(p.extract_text() or "" for p in pdf.pages)
    elif ext in (".xlsx", ".xls"):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        rows = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                rows.append(" | ".join(str(c) for c in row if c is not None))
        return "\n".join(rows)
    elif ext in (".docx", ".doc"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"Unsupported file type: {ext}")
